"""Característica general 5 — Reportes personalizables y exportables.

Lo que la consigna de la materia pide, punto por punto, y una prueba por cada
cosa que podría romperse en silencio:

- el usuario elige **qué columnas**, **qué criterios de selección** y **qué
  orden**;
- hay una **interfaz previa** para filtrar antes de generar — acá es
  ``/datasets/``, que le dice al formulario qué filtros existen y de qué tipo;
- todo reporte se **exporta a Excel, HTML, correo y PDF**.

Las que más importan son las dos que no salen del enunciado: que un reporte no
sea una puerta lateral para leer lo que la pantalla niega, y que un filtro que
no se pudo aplicar **falle** en vez de devolver la tabla entera.
"""

import io
import json

import pytest
from django.core import mail
from django.urls import reverse
from rest_framework.test import APIClient

from accounts.models import AuditLog
from accounts.tokens import tokens_for_user
from audit.actions import Action
from patients.models import Patient
from reporting.models import SavedReport
from tenancy.context import tenant_context

from .conftest import dar_rol

pytestmark = pytest.mark.django_db

# Quien arma reportes de pacientes necesita las dos cosas: poder usar el
# constructor y poder leer pacientes.
PERMISOS_ANALISTA = [
    "reporting.report.run",
    "reporting.report.save",
    "patients.patient.read",
]


@pytest.fixture
def api_client():
    return APIClient()


def autenticar(api_client, user):
    api_client.credentials(
        HTTP_AUTHORIZATION=f"Bearer {tokens_for_user(user)['access']}",
    )
    return api_client


@pytest.fixture
def analista_a(db, org_a, user_a):
    dar_rol(user_a, org_a, "analista", "Analista", PERMISOS_ANALISTA)
    return user_a


@pytest.fixture
def pacientes_a(db, org_a):
    """Seis pacientes con datos distinguibles, para poder probar los filtros.

    Son ficticios: el repositorio es público y la regla 7 del reparto no admite
    datos de personas reales ni siquiera en las pruebas.
    """
    with tenant_context(org_a.id):
        return [
            Patient.objects.create(
                organization=org_a, first_name=nombre, last_name=apellido,
                document_number=documento, sex=sexo, is_active=activo,
            )
            for nombre, apellido, documento, sexo, activo in [
                ("Ana", "Álvarez", "9001", "F", True),
                ("Bruno", "Bustos", "9002", "M", True),
                ("Carla", "Cárdenas", "9003", "F", True),
                ("Diego", "Duarte", "9004", "M", False),
                ("Elena", "Escobar", "9005", "F", False),
                ("Fabio", "Flores", "9006", "M", True),
            ]
        ]


def correr(api_client, **cuerpo):
    return api_client.post(reverse("reporting:run"), cuerpo, format="json")


# --------------------------------------------------------------------------
#  La interfaz previa: qué se puede reportar y con qué filtros
# --------------------------------------------------------------------------
def test_el_catalogo_describe_columnas_filtros_y_operadores(
    api_client, analista_a,
):
    """Sin esto el frontend tendría que conocer el modelo de datos.

    Es lo que permite que el constructor se dibuje solo: tipos para elegir el
    control, operadores para armar la comparación, opciones para el
    desplegable.
    """
    respuesta = autenticar(api_client, analista_a).get(
        reverse("reporting:datasets"),
    )
    assert respuesta.status_code == 200

    conjuntos = {c["code"]: c for c in respuesta.data["datasets"]}
    assert "patients" in conjuntos

    pacientes = conjuntos["patients"]
    columnas = {c["code"]: c for c in pacientes["columns"]}
    assert columnas["birth_date"]["kind"] == "date"
    assert columnas["is_active"]["kind"] == "boolean"
    # El sexo llega con sus opciones legibles, no con los códigos pelados.
    assert {"value": "F", "label": "Femenino"} in columnas["sex"]["choices"]

    filtros = {f["code"]: f for f in pacientes["filters"]}
    assert "contains" in filtros["last_name"]["operators"]
    assert "gte" in filtros["birth_date"]["operators"]

    assert set(respuesta.data["formats"]) == {"csv", "xlsx", "html", "pdf"}


def test_el_catalogo_solo_ofrece_lo_que_la_persona_puede_leer(
    api_client, org_a, user_a,
):
    """Un desplegable con seis opciones de las que cinco dan 403 no sirve."""
    dar_rol(user_a, org_a, "limitado", "Limitado",
            ["reporting.report.run", "catalog.branch.read"])

    respuesta = autenticar(api_client, user_a).get(
        reverse("reporting:datasets"),
    )
    codigos = {c["code"] for c in respuesta.data["datasets"]}
    assert codigos == {"branches"}


# --------------------------------------------------------------------------
#  Columnas, criterios y orden: los tres que pide el enunciado
# --------------------------------------------------------------------------
def test_el_usuario_elige_las_columnas(api_client, analista_a, pacientes_a):
    respuesta = correr(
        autenticar(api_client, analista_a),
        dataset="patients", columns=["last_name", "document_number"],
    )
    assert respuesta.status_code == 200
    assert [c["label"] for c in respuesta.data["columns"]] == [
        "Apellidos", "Documento",
    ]
    assert all(len(fila) == 2 for fila in respuesta.data["rows"])


def test_el_usuario_elige_los_criterios_de_seleccion(
    api_client, analista_a, pacientes_a,
):
    respuesta = correr(
        autenticar(api_client, analista_a),
        dataset="patients", columns=["last_name"],
        filters=[{"field": "is_active", "operator": "eq", "value": False}],
    )
    assert respuesta.status_code == 200
    assert sorted(fila[0] for fila in respuesta.data["rows"]) == [
        "Duarte", "Escobar",
    ]


def test_el_usuario_elige_el_orden(api_client, analista_a, pacientes_a):
    respuesta = correr(
        autenticar(api_client, analista_a),
        dataset="patients", columns=["last_name"], order_by=["-last_name"],
    )
    assert [fila[0] for fila in respuesta.data["rows"]][:2] == [
        "Flores", "Escobar",
    ]


def test_un_codigo_de_choices_sale_con_su_etiqueta(
    api_client, analista_a, pacientes_a,
):
    """En el archivo tiene que decir «Femenino», no «F».

    Es la diferencia entre un reporte que se entrega y uno que hay que
    traducir a mano.
    """
    respuesta = correr(
        autenticar(api_client, analista_a),
        dataset="patients", columns=["sex"],
        filters=[{"field": "document_number", "value": "9001"}],
    )
    assert respuesta.data["rows"] == [["Femenino"]]


def test_un_booleano_sale_en_castellano(api_client, analista_a, pacientes_a):
    respuesta = correr(
        autenticar(api_client, analista_a),
        dataset="patients", columns=["is_active"],
        filters=[{"field": "document_number", "value": "9004"}],
    )
    assert respuesta.data["rows"] == [["No"]]


# --------------------------------------------------------------------------
#  Lo que tiene que fallar, y fallar ruidosamente
# --------------------------------------------------------------------------
def test_una_columna_que_no_esta_en_el_catalogo_se_rechaza(
    api_client, analista_a, pacientes_a,
):
    """La lista blanca es lo único que impide llegar a `user__password`.

    Sin ella, cualquier ruta del ORM sería una columna válida.
    """
    respuesta = correr(
        autenticar(api_client, analista_a),
        dataset="patients", columns=["user__password"],
    )
    assert respuesta.status_code == 400
    assert respuesta.data["code"] == "columna_desconocida"


def test_no_se_puede_ordenar_por_algo_que_no_es_columna(
    api_client, analista_a, pacientes_a,
):
    """Ordenar por un campo que no se muestra sigue filtrando por él.

    Con `order_by` libre se puede adivinar un valor oculto letra por letra
    mirando cómo cambia el orden, sin que ese campo aparezca nunca.
    """
    respuesta = correr(
        autenticar(api_client, analista_a),
        dataset="patients", columns=["last_name"], order_by=["user__password"],
    )
    assert respuesta.status_code == 400
    assert respuesta.data["code"] == "orden_desconocido"


def test_un_filtro_desconocido_no_se_ignora_en_silencio(
    api_client, analista_a, pacientes_a,
):
    """El defecto clásico de un filtro dinámico.

    Si un criterio mal escrito no filtra nada, el reporte sale con la tabla
    entera y quien lo lee concluye que ése es el dato. Tiene que fallar.
    """
    respuesta = correr(
        autenticar(api_client, analista_a),
        dataset="patients", columns=["last_name"],
        filters=[{"field": "apellido", "value": "Álvarez"}],
    )
    assert respuesta.status_code == 400
    assert respuesta.data["code"] == "filtro_desconocido"
    # El mensaje dice por qué se puede filtrar: el error tiene que enseñar.
    assert "last_name" in respuesta.data["detail"]


def test_un_reporte_no_es_una_puerta_lateral(api_client, org_a, user_a):
    """Quien no puede leer la bitácora tampoco puede exportarla.

    Es la razón de que el permiso del conjunto se comprueba además del de
    reportes: sin esto, `reporting.report.run` sería un permiso para leerlo
    todo.
    """
    dar_rol(user_a, org_a, "analista", "Analista", PERMISOS_ANALISTA)

    respuesta = correr(
        autenticar(api_client, user_a),
        dataset="audit", columns=["action"],
    )
    assert respuesta.status_code == 403
    assert respuesta.data["code"] == "sin_permiso_sobre_el_conjunto"


def test_sin_permiso_de_reportes_no_se_entra_aunque_pueda_leer_el_dato(
    api_client, org_a, user_a,
):
    """La otra mitad: poder ver pacientes no es poder exportarlos.

    Mirar una ficha por pantalla y bajarse el padrón entero en un Excel no son
    la misma acción, y la organización tiene que poder separarlas.
    """
    dar_rol(user_a, org_a, "recepcion", "Recepción", ["patients.patient.read"])

    respuesta = correr(
        autenticar(api_client, user_a), dataset="patients",
    )
    assert respuesta.status_code == 403


# --------------------------------------------------------------------------
#  Exportación: los cuatro formatos del enunciado
# --------------------------------------------------------------------------
@pytest.mark.parametrize("formato, tipo, firma", [
    ("csv", "text/csv", b"\xef\xbb\xbf"),          # BOM: Excel en español
    ("xlsx", "spreadsheetml", b"PK"),              # un .xlsx es un zip
    ("html", "text/html", b"<!doctype html>"),
    ("pdf", "application/pdf", b"%PDF"),
])
def test_se_exporta_a_los_cuatro_formatos(
    api_client, analista_a, pacientes_a, formato, tipo, firma,
):
    """Que el archivo sea de verdad del formato que dice.

    Se comprueba la firma de los primeros bytes y no sólo el `Content-Type`:
    una cabecera correcta sobre un cuerpo vacío es el fallo que se descubre
    recién cuando alguien intenta abrir el archivo.
    """
    respuesta = correr(
        autenticar(api_client, analista_a),
        dataset="patients", columns=["last_name", "document_number"],
        format=formato, title="Padrón de prueba",
    )
    assert respuesta.status_code == 200
    assert tipo in respuesta["Content-Type"]
    assert respuesta["Content-Disposition"].startswith("attachment;")
    assert f".{formato}" in respuesta["Content-Disposition"]

    contenido = respuesta.content
    assert contenido[: len(firma)].lower() == firma.lower()
    assert len(contenido) > 100


def test_el_excel_se_abre_y_trae_los_datos(
    api_client, analista_a, pacientes_a,
):
    """No alcanza con que empiece por «PK»: hay que poder leerlo de vuelta."""
    from openpyxl import load_workbook

    respuesta = correr(
        autenticar(api_client, analista_a),
        dataset="patients", columns=["last_name", "sex"],
        filters=[{"field": "sex", "value": "F"}],
        order_by=["last_name"], format="xlsx", title="Pacientes mujeres",
    )
    hoja = load_workbook(io.BytesIO(respuesta.content)).active

    valores = [
        [celda.value for celda in fila]
        for fila in hoja.iter_rows(min_row=1, max_col=2)
    ]
    planas = [v for fila in valores for v in fila if v]

    assert "Pacientes mujeres" in planas[0]
    assert "Apellidos" in planas
    assert "Álvarez" in planas
    assert "Femenino" in planas
    # El filtro se respetó también en el archivo, no sólo en la vista previa.
    assert "Bustos" not in planas


def test_el_csv_lleva_bom_y_punto_y_coma(api_client, analista_a, pacientes_a):
    """Las dos decisiones que hacen que Excel en español lo abra bien.

    Sin BOM, «Álvarez» sale «Ãlvarez». Sin punto y coma, todo el reporte cae
    en la primera columna.
    """
    respuesta = correr(
        autenticar(api_client, analista_a),
        dataset="patients", columns=["last_name", "document_number"],
        format="csv",
    )
    texto = respuesta.content.decode("utf-8-sig")
    assert respuesta.content.startswith(b"\xef\xbb\xbf")
    assert texto.splitlines()[0] == "Apellidos;Documento"
    assert "Álvarez;9001" in texto


def test_el_archivo_dice_con_que_criterios_salio(
    api_client, analista_a, pacientes_a,
):
    """Un PDF reenviado por correo no recuerda sus filtros si no los lleva.

    Dos exportaciones del mismo reporte con criterios distintos serían
    indistinguibles, que es como se toma una decisión sobre el reporte
    equivocado.
    """
    respuesta = correr(
        autenticar(api_client, analista_a),
        dataset="patients", columns=["last_name"],
        filters=[{"field": "sex", "value": "F"}], format="html",
    )
    html = respuesta.content.decode("utf-8")
    assert "Criterios" in html
    assert "Sexo: Femenino" in html


# --------------------------------------------------------------------------
#  El correo
# --------------------------------------------------------------------------
def test_el_reporte_se_manda_por_correo_adjunto(
    api_client, analista_a, pacientes_a,
):
    respuesta = correr(
        autenticar(api_client, analista_a),
        dataset="patients", columns=["last_name"], format="pdf",
        title="Padrón", recipients=[analista_a.email],
    )
    assert respuesta.status_code == 200
    assert respuesta.data["sent"] is True

    assert len(mail.outbox) == 1
    mensaje = mail.outbox[0]
    assert mensaje.to == [analista_a.email.lower()]
    nombre, contenido, mime = mensaje.attachments[0]
    assert nombre.endswith(".pdf")
    assert contenido[:4] == b"%PDF"
    assert mime == "application/pdf"


def test_no_se_puede_mandar_el_reporte_a_una_casilla_ajena(
    api_client, analista_a, pacientes_a,
):
    """Sin esto el endpoint es un relé de correo abierto.

    Cualquiera con una cuenta podría mandar lo que quisiera a cualquier
    dirección, firmado por el centro médico.
    """
    respuesta = correr(
        autenticar(api_client, analista_a),
        dataset="patients", columns=["last_name"], format="pdf",
        recipients=["cualquiera@internet.test"],
    )
    assert respuesta.status_code == 400
    assert respuesta.data["code"] == "destinatario_no_permitido"
    assert mail.outbox == []


def test_pedir_correo_sin_formato_de_archivo_se_rechaza(
    api_client, analista_a, pacientes_a,
):
    """Adjuntar un `.json` que nadie puede abrir no es enviar un reporte."""
    respuesta = correr(
        autenticar(api_client, analista_a),
        dataset="patients", recipients=[analista_a.email],
    )
    assert respuesta.status_code == 400
    assert "format" in respuesta.data


# --------------------------------------------------------------------------
#  Reportes guardados
# --------------------------------------------------------------------------
def test_se_guarda_un_reporte_y_se_vuelve_a_correr(
    api_client, analista_a, pacientes_a,
):
    """Es lo que convierte «armar un reporte» en «tener un reporte»."""
    cliente = autenticar(api_client, analista_a)

    creado = cliente.post(reverse("reporting:saved_report-list"), {
        "name": "Pacientes activos",
        "description": "Los que siguen en el padrón",
        "definition": {
            "dataset": "patients",
            "columns": ["last_name", "document_number"],
            "filters": [{"field": "is_active", "operator": "eq", "value": True}],
            "order_by": ["last_name"],
        },
    }, format="json")
    assert creado.status_code == 201
    assert creado.data["dataset"] == "patients"
    assert creado.data["dataset_label"] == "Pacientes"

    corrido = cliente.post(
        reverse("reporting:saved_report-run", args=[creado.data["id"]]),
        {}, format="json",
    )
    assert corrido.status_code == 200
    assert corrido.data["title"] == "Pacientes activos"
    assert [fila[0] for fila in corrido.data["rows"]] == [
        "Álvarez", "Bustos", "Cárdenas", "Flores",
    ]


def test_al_correr_uno_guardado_se_le_pueden_cambiar_los_filtros(
    api_client, analista_a, pacientes_a,
):
    """El caso normal: el mismo reporte, otro período.

    Guardar un reporte por cada mes sería absurdo, así que el cuerpo de la
    llamada se superpone a lo guardado.
    """
    cliente = autenticar(api_client, analista_a)
    creado = cliente.post(reverse("reporting:saved_report-list"), {
        "name": "Por sexo",
        "definition": {
            "dataset": "patients", "columns": ["last_name"],
            "filters": [{"field": "sex", "value": "F"}],
        },
    }, format="json")

    corrido = cliente.post(
        reverse("reporting:saved_report-run", args=[creado.data["id"]]),
        {"filters": [{"field": "sex", "value": "M"}]}, format="json",
    )
    assert sorted(fila[0] for fila in corrido.data["rows"]) == [
        "Bustos", "Duarte", "Flores",
    ]


def test_compartir_un_reporte_exige_su_propio_permiso(
    api_client, analista_a,
):
    """Compartir cambia lo que ve el resto de la organización.

    El analista puede guardar lo suyo; ponerlo a la vista de todos es una
    decisión de quien administra.
    """
    respuesta = autenticar(api_client, analista_a).post(
        reverse("reporting:saved_report-list"), {
            "name": "Para todos",
            "is_shared": True,
            "definition": {"dataset": "patients", "columns": ["last_name"]},
        }, format="json")
    assert respuesta.status_code == 403


def test_un_reporte_compartido_no_presta_permisos(
    api_client, org_a, user_a, user_b,
):
    """Quien lo corre necesita el permiso del dato, no el de quien lo guardó.

    Si no fuera así, compartir un reporte sería la forma de regalar acceso a
    la bitácora sin pasar por la pantalla de roles.
    """
    admin = dar_rol(user_a, org_a, "admin_rep", "Admin",
                    PERMISOS_ANALISTA + ["reporting.report.share",
                                         "audit.log.read"])
    assert admin is not None

    creado = autenticar(api_client, user_a).post(
        reverse("reporting:saved_report-list"), {
            "name": "Bitácora compartida", "is_shared": True,
            "definition": {"dataset": "audit", "columns": ["action"]},
        }, format="json")
    assert creado.status_code == 201

    # Otro usuario de la MISMA organización, sin permiso sobre la bitácora.
    with tenant_context(org_a.id):
        from accounts.models import User
        otro = User.objects.create_user(
            organization=org_a, email="otro@a.test",
            password="clave-de-prueba-1", first_name="Otro", last_name="Uno",
        )
    dar_rol(otro, org_a, "solo_reportes", "Sólo reportes",
            ["reporting.report.run"])

    cliente = APIClient()
    respuesta = autenticar(cliente, otro).post(
        reverse("reporting:saved_report-run", args=[creado.data["id"]]),
        {}, format="json",
    )
    assert respuesta.status_code == 403


def test_solo_el_dueno_edita_su_reporte(api_client, org_a, user_a):
    """Un reporte compartido que otro reescribe sin avisar deja de ser
    confiable para quien lo corre todos los lunes."""
    dar_rol(user_a, org_a, "admin_rep", "Admin",
            PERMISOS_ANALISTA + ["reporting.report.share"])

    creado = autenticar(api_client, user_a).post(
        reverse("reporting:saved_report-list"), {
            "name": "Mío", "is_shared": True,
            "definition": {"dataset": "patients", "columns": ["last_name"]},
        }, format="json")

    with tenant_context(org_a.id):
        from accounts.models import User
        otro = User.objects.create_user(
            organization=org_a, email="otro2@a.test",
            password="clave-de-prueba-1", first_name="Otro", last_name="Dos",
        )
    dar_rol(otro, org_a, "analista2", "Analista", PERMISOS_ANALISTA)

    cliente = APIClient()
    respuesta = autenticar(cliente, otro).patch(
        reverse("reporting:saved_report-detail", args=[creado.data["id"]]),
        {"name": "Ya no"}, format="json",
    )
    assert respuesta.status_code == 403

    # Pero sí lo ve, porque está compartido.
    listado = autenticar(cliente, otro).get(
        reverse("reporting:saved_report-list"),
    )
    nombres = [r["name"] for r in listado.data["results"]]
    assert "Mío" in nombres


# --------------------------------------------------------------------------
#  El cruce con la característica 3: toda exportación deja rastro
# --------------------------------------------------------------------------
def test_generar_un_reporte_deja_asiento_en_la_bitacora(
    api_client, org_a, analista_a, pacientes_a,
):
    """Exportar el padrón es sacar datos del sistema.

    Es la única lectura masiva que tiene la plataforma, y es exactamente lo
    que una bitácora existe para poder contar después.
    """
    correr(
        autenticar(api_client, analista_a),
        dataset="patients", columns=["last_name"], format="xlsx",
        filters=[{"field": "is_active", "value": True}],
    )

    with tenant_context(org_a.id):
        asiento = AuditLog.objects.filter(action=Action.REPORT_RUN).first()

    assert asiento is not None
    assert asiento.user_id == analista_a.id
    assert asiento.detail["dataset"] == "patients"
    assert asiento.detail["format"] == "xlsx"
    assert asiento.detail["rows"] == 4
    # Los filtros quedan; las filas NO. Una bitácora que copia el contenido de
    # cada reporte es una segunda base con los mismos datos y ninguna
    # protección.
    assert asiento.detail["filters"] == [
        {"field": "is_active", "value": True},
    ]
    assert "Álvarez" not in json.dumps(asiento.detail)


def test_mandar_un_reporte_por_correo_se_audita_aparte(
    api_client, org_a, analista_a, pacientes_a,
):
    """Mandar el padrón a una casilla no es lo mismo que mirarlo en pantalla,
    y quien audita quiere poder filtrar exactamente eso."""
    correr(
        autenticar(api_client, analista_a),
        dataset="patients", columns=["last_name"], format="pdf",
        recipients=[analista_a.email],
    )

    with tenant_context(org_a.id):
        asiento = AuditLog.objects.filter(action=Action.REPORT_EMAIL).first()

    assert asiento is not None
    assert asiento.detail["recipients"] == [analista_a.email.lower()]


# --------------------------------------------------------------------------
#  Aislamiento por HTTP
# --------------------------------------------------------------------------
def test_un_reporte_solo_alcanza_los_datos_de_su_organizacion(
    api_client, org_a, org_b, analista_a, pacientes_a,
):
    """El corazón del proyecto, aplicado a los reportes.

    Un constructor de reportes es justamente donde un fallo de aislamiento se
    convierte en un archivo con el padrón del vecino adentro.
    """
    with tenant_context(org_b.id):
        Patient.objects.create(
            organization=org_b, first_name="Zulma", last_name="Zapata",
            document_number="9999",
        )

    respuesta = correr(
        autenticar(api_client, analista_a),
        dataset="patients", columns=["last_name"],
    )
    apellidos = [fila[0] for fila in respuesta.data["rows"]]
    assert "Zapata" not in apellidos
    assert len(apellidos) == 6


def test_el_superadministrador_no_reporta_datos_de_ningun_inquilino(
    api_client, platform_admin,
):
    """Su alcance no incluye los datos internos de ninguna organización, y un
    reporte no es la excepción."""
    respuesta = correr(
        autenticar(api_client, platform_admin),
        dataset="patients", columns=["last_name"],
    )
    # No tiene el permiso de reportes: su rol sólo lleva los del módulo
    # `platform`. Falla antes de llegar a los datos, que es lo correcto.
    assert respuesta.status_code == 403
