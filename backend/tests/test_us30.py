"""US-30 — Reportes operativos: los conjuntos de fichas y de pagos.

La exportación a PDF y Excel ya existía (característica general 5); acá se
prueba lo que esta historia agrega: que fichas y pagos se pueden pedir, que
sólo los ve la administración y que **los números coinciden con el panel de
US-29** (`docs/sprints/sprint-3/us-29-indicadores.md`).
"""

import datetime as dt
import io
from decimal import Decimal
from zoneinfo import ZoneInfo

import pytest
from django.urls import reverse
from django.utils import timezone

from payments.models import Payment
from tenancy.context import tenant_context

from .conftest import dar_rol
from .fichas_comunes import (  # noqa: F401  (fixtures)
    agenda_a,
    api_client,
    autenticar,
    crear_ficha,
    paciente_a,
)

pytestmark = pytest.mark.django_db

LA_PAZ = ZoneInfo("America/La_Paz")

# Los estados que ocupan un espacio: los que el panel cuenta como demanda.
OCUPAN_ESPACIO = ["pending_payment", "confirmed", "attended", "no_show"]


@pytest.fixture
def admin_a(paciente_a, org_a):
    """`user_a` con el permiso del panel y el del constructor de reportes."""
    dar_rol(paciente_a, org_a, "admin-operativo", "Administración",
            ["reporting.kpi.read", "reporting.report.run"])
    return paciente_a


def hace(dias):
    """Una hora redonda `dias` atrás (o adelante si es negativo)."""
    base = timezone.now() - dt.timedelta(days=dias)
    return base.replace(minute=0, second=0, microsecond=0)


def rango(dias_atras=10, dias_adelante=10):
    hoy = timezone.now().astimezone(LA_PAZ).date()
    return {
        "from": (hoy - dt.timedelta(days=dias_atras)).isoformat(),
        "to": (hoy + dt.timedelta(days=dias_adelante)).isoformat(),
    }


def ficha(org, paciente, practitioner, branch, schedule, dias, **extra):
    return crear_ficha(org, paciente, practitioner, branch, schedule,
                       starts_at=hace(dias), **extra)


def pagar(org, ficha_, usuario, monto, *, estado="succeeded", moneda="BOB",
          medio="simulated"):
    with tenant_context(org.id):
        return Payment.objects.create(
            organization=org, appointment=ficha_, created_by=usuario,
            amount=Decimal(monto), currency=moneda, provider=medio,
            status=estado,
        )


def correr(client, **cuerpo):
    return client.post(reverse("reporting:run"), cuerpo, format="json")


def en_el_rango(**extra):
    """Los dos criterios de fecha que equivalen al `from`/`to` del panel."""
    r = rango()
    return [
        {"field": "starts_at", "operator": "gte", "value": r["from"]},
        {"field": "starts_at", "operator": "lte", "value": r["to"]},
        *extra.get("mas", []),
    ]


# ---------------------------------------------------------------------------
#  Quién los ve
# ---------------------------------------------------------------------------

def test_la_administracion_ve_fichas_y_pagos_en_el_catalogo(api_client, admin_a):
    respuesta = autenticar(api_client, admin_a).get(reverse("reporting:datasets"))

    conjuntos = {c["code"]: c for c in respuesta.data["datasets"]}
    assert {"appointments", "payments"} <= set(conjuntos)

    columnas = {c["code"]: c for c in conjuntos["appointments"]["columns"]}
    assert {"value": "absent", "label": "Ausente"} in columnas["attendance"]["choices"]
    assert {"value": "confirmed", "label": "Confirmada"} in columnas["status"]["choices"]
    medios = {c["code"]: c for c in conjuntos["payments"]["columns"]}["provider"]
    assert {"value": "cash", "label": "Efectivo en ventanilla"} in medios["choices"]


def test_recepcion_con_reportes_no_ve_fichas_ni_pagos(api_client, org_a, user_a):
    """Recepción arma listados y lee fichas, pero los ingresos no son suyos."""
    dar_rol(user_a, org_a, "recepcion-reportes", "Recepción",
            ["reporting.report.run", "appointments.appointment.read"])
    cliente = autenticar(api_client, user_a)

    codigos = {c["code"] for c in cliente.get(reverse("reporting:datasets")).data["datasets"]}
    assert "appointments" not in codigos
    assert "payments" not in codigos

    for conjunto in ("appointments", "payments"):
        assert correr(cliente, dataset=conjunto).status_code == 403


def test_un_paciente_no_se_lleva_las_fichas_del_centro(api_client, org_a, paciente_a):
    """`appointments.appointment.read` lo tienen los pacientes: no abre el reporte."""
    dar_rol(paciente_a, org_a, "paciente-curioso", "Paciente",
            ["reporting.report.run", "appointments.appointment.read"])

    respuesta = correr(autenticar(api_client, paciente_a), dataset="appointments")

    assert respuesta.status_code == 403


# ---------------------------------------------------------------------------
#  Fichas: los mismos números que el panel
# ---------------------------------------------------------------------------

def _fichas_del_panel(org, admin, practitioner, centro, norte, agenda):
    """Las mismas fichas del test de inasistencia del panel, más una de otra sede."""
    ficha(org, admin, practitioner, centro, agenda, 9, status="attended")
    ficha(org, admin, practitioner, centro, agenda, 8, status="no_show")
    # Confirmada que ya pasó y nunca hizo check-in: el panel la cuenta ausente.
    ficha(org, admin, practitioner, centro, agenda, 7, status="confirmed")
    ficha(org, admin, practitioner, centro, agenda, 6, status="cancelled")
    ficha(org, admin, practitioner, centro, agenda, 5, status="rescheduled")
    ficha(org, admin, practitioner, centro, agenda, 4, status="expired")
    ficha(org, admin, practitioner, norte, agenda, 3, status="pending_payment")
    ficha(org, admin, practitioner, norte, agenda, -2, status="confirmed")


def test_las_ausentes_del_reporte_son_las_del_panel(
    api_client, admin_a, org_a, practitioner_a, branches_a, agenda_a,
):
    _fichas_del_panel(org_a, admin_a, practitioner_a, branches_a["centro"],
                      branches_a["norte"], agenda_a)
    cliente = autenticar(api_client, admin_a)

    panel = cliente.get(reverse("reporting:kpis"), rango()).data["no_show"]
    assert (panel["expected"], panel["missed"]) == (3, 2)

    def contar(asistencia):
        respuesta = correr(
            cliente, dataset="appointments", columns=["status", "attendance"],
            filters=en_el_rango(mas=[
                {"field": "attendance", "operator": "eq", "value": asistencia},
            ]),
        )
        assert respuesta.status_code == 200
        return respuesta.data["rows"]

    ausentes = contar("absent")
    assert len(ausentes) == panel["missed"]
    # La confirmada vencida sale «Ausente» aunque su estado siga siendo «Confirmada».
    assert sorted(fila[0] for fila in ausentes) == ["Ausente", "Confirmada"]
    assert {fila[1] for fila in ausentes} == {"Ausente"}

    asistieron = contar("attended")
    assert len(ausentes) + len(asistieron) == panel["expected"]


def test_las_fichas_del_reporte_cuadran_con_la_demanda_del_panel(
    api_client, admin_a, org_a, practitioner_a, branches_a, agenda_a,
):
    _fichas_del_panel(org_a, admin_a, practitioner_a, branches_a["centro"],
                      branches_a["norte"], agenda_a)
    cliente = autenticar(api_client, admin_a)

    panel = cliente.get(reverse("reporting:kpis"), rango()).data["demand"]

    respuesta = correr(
        cliente, dataset="appointments", columns=["branch", "status"],
        filters=en_el_rango(mas=[
            {"field": "status", "operator": "in", "value": OCUPAN_ESPACIO},
        ]),
    )
    filas = respuesta.data["rows"]

    assert len(filas) == panel["total"]
    por_sucursal = {}
    for sucursal, _estado in filas:
        por_sucursal[sucursal] = por_sucursal.get(sucursal, 0) + 1
    assert por_sucursal == {s["name"]: s["total"] for s in panel["by_branch"]}


def test_el_rango_de_fechas_deja_afuera_lo_que_no_cae_en_el(
    api_client, admin_a, org_a, practitioner_a, branches_a, agenda_a,
):
    centro = branches_a["centro"]
    ficha(org_a, admin_a, practitioner_a, centro, agenda_a, 40, status="attended")
    ficha(org_a, admin_a, practitioner_a, centro, agenda_a, 5, status="attended")

    respuesta = correr(
        autenticar(api_client, admin_a), dataset="appointments",
        columns=["starts_at"], filters=en_el_rango(),
    )

    assert len(respuesta.data["rows"]) == 1


def test_con_pocas_columnas_sale_una_fila_por_ficha(
    api_client, admin_a, org_a, practitioner_a, branches_a, agenda_a,
):
    """«Fichas por estado» se saca contando filas: no pueden colapsar."""
    centro = branches_a["centro"]
    for dias in (9, 8, 7):
        ficha(org_a, admin_a, practitioner_a, centro, agenda_a, dias, status="attended")
    ficha(org_a, admin_a, practitioner_a, centro, agenda_a, 6, status="cancelled")

    respuesta = correr(
        autenticar(api_client, admin_a), dataset="appointments",
        columns=["status"], filters=en_el_rango(),
    )

    estados = [fila[0] for fila in respuesta.data["rows"]]
    assert estados.count("Atendida") == 3
    assert estados.count("Cancelada") == 1


def test_el_reporte_de_fichas_se_puede_filtrar_por_sucursal_y_ordenar(
    api_client, admin_a, org_a, practitioner_a, branches_a, agenda_a,
):
    ficha(org_a, admin_a, practitioner_a, branches_a["centro"], agenda_a, 9, status="attended")
    ficha(org_a, admin_a, practitioner_a, branches_a["norte"], agenda_a, 8, status="attended")
    ficha(org_a, admin_a, practitioner_a, branches_a["norte"], agenda_a, 7, status="cancelled")

    respuesta = correr(
        autenticar(api_client, admin_a), dataset="appointments",
        columns=["branch", "status", "starts_at"],
        filters=[{"field": "branch", "operator": "eq", "value": "Sede Norte"}],
        order_by=["starts_at"],
    )

    assert [fila[:2] for fila in respuesta.data["rows"]] == [
        ["Sede Norte", "Atendida"], ["Sede Norte", "Cancelada"],
    ]


# ---------------------------------------------------------------------------
#  Pagos: los mismos ingresos que el panel
# ---------------------------------------------------------------------------

def _pagos_del_panel(org, admin, practitioner, centro, norte, agenda):
    f1 = ficha(org, admin, practitioner, centro, agenda, 9, status="attended")
    f2 = ficha(org, admin, practitioner, norte, agenda, 8, status="attended")
    f3 = ficha(org, admin, practitioner, centro, agenda, 7, status="cancelled")
    f4 = ficha(org, admin, practitioner, centro, agenda, 6, status="attended")
    f5 = ficha(org, admin, practitioner, centro, agenda, 5, status="attended")
    pagar(org, f1, admin, "100.00")
    pagar(org, f2, admin, "50.50", medio="cash")
    pagar(org, f3, admin, "80.00", estado="refunded")  # devuelto: no suma
    pagar(org, f4, admin, "40.00", estado="pending")   # sin cobrar: no suma
    pagar(org, f5, admin, "10.00", moneda="USD")


def test_los_ingresos_del_reporte_son_los_del_panel(
    api_client, admin_a, org_a, practitioner_a, branches_a, agenda_a,
):
    _pagos_del_panel(org_a, admin_a, practitioner_a, branches_a["centro"],
                     branches_a["norte"], agenda_a)
    cliente = autenticar(api_client, admin_a)

    panel = cliente.get(reverse("reporting:kpis"), rango()).data["revenue"]

    respuesta = correr(
        cliente, dataset="payments", columns=["branch", "amount", "currency"],
        filters=[
            {"field": "appointment_starts_at", "operator": "gte", "value": rango()["from"]},
            {"field": "appointment_starts_at", "operator": "lte", "value": rango()["to"]},
            {"field": "status", "operator": "eq", "value": "succeeded"},
        ],
    )
    assert respuesta.status_code == 200

    por_moneda, por_sucursal = {}, {}
    for sucursal, importe, moneda in respuesta.data["rows"]:
        importe = Decimal(str(importe))
        por_moneda[moneda] = por_moneda.get(moneda, 0) + importe
        por_sucursal[(sucursal, moneda)] = por_sucursal.get((sucursal, moneda), 0) + importe

    assert {m: str(t) for m, t in por_moneda.items()} == {
        f["currency"]: f["total"] for f in panel["by_currency"]
    }
    assert {k: str(t) for k, t in por_sucursal.items()} == {
        (f["name"], f["currency"]): f["total"] for f in panel["by_branch"]
    }


def test_el_pago_en_ventanilla_sale_con_su_nombre(
    api_client, admin_a, org_a, practitioner_a, branches_a, agenda_a,
):
    f = ficha(org_a, admin_a, practitioner_a, branches_a["centro"], agenda_a, 5, status="confirmed")
    pagar(org_a, f, admin_a, "30.00", medio="cash")

    respuesta = correr(
        autenticar(api_client, admin_a), dataset="payments",
        columns=["provider", "status"],
        filters=[{"field": "provider", "operator": "eq", "value": "cash"}],
    )

    assert respuesta.data["rows"] == [["Efectivo en ventanilla", "Pagado"]]


# ---------------------------------------------------------------------------
#  Exportación a PDF y Excel
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("conjunto", ["appointments", "payments"])
@pytest.mark.parametrize("formato, firma", [("pdf", b"%PDF"), ("xlsx", b"PK")])
def test_fichas_y_pagos_se_exportan_a_pdf_y_a_excel(
    api_client, admin_a, org_a, practitioner_a, branches_a, agenda_a,
    conjunto, formato, firma,
):
    f = ficha(org_a, admin_a, practitioner_a, branches_a["centro"], agenda_a, 5, status="attended")
    pagar(org_a, f, admin_a, "100.00")

    respuesta = correr(
        autenticar(api_client, admin_a), dataset=conjunto, format=formato,
        title="Reporte para la dirección",
    )

    assert respuesta.status_code == 200
    assert respuesta.content[: len(firma)] == firma
    assert f".{formato}" in respuesta["Content-Disposition"]
    assert respuesta["X-Report-Rows"] == "1"


def test_el_importe_llega_al_excel_como_numero_que_se_puede_sumar(
    api_client, admin_a, org_a, practitioner_a, branches_a, agenda_a,
):
    from openpyxl import load_workbook

    for dias, monto in ((9, "100.00"), (8, "50.50")):
        f = ficha(org_a, admin_a, practitioner_a, branches_a["centro"], agenda_a, dias, status="attended")
        pagar(org_a, f, admin_a, monto)

    respuesta = correr(
        autenticar(api_client, admin_a), dataset="payments",
        columns=["amount", "currency"], format="xlsx", title="Ingresos",
    )
    hoja = load_workbook(io.BytesIO(respuesta.content)).active
    importes = [
        celda.value for fila in hoja.iter_rows()
        for celda in fila if isinstance(celda.value, (int, float))
    ]

    assert sorted(importes) == [50.5, 100.0]
    assert sum(importes) == pytest.approx(150.5)


# ---------------------------------------------------------------------------
#  Aislamiento entre organizaciones (RNF-08)
# ---------------------------------------------------------------------------

@pytest.mark.isolation
def test_el_reporte_solo_trae_los_datos_de_la_organizacion(
    api_client, admin_a, org_a, practitioner_a, branches_a, agenda_a,
    org_b, user_b, branch_b,
):
    from catalog.models import Practitioner, PractitionerBranch
    from scheduling.models import Schedule

    from .fichas_comunes import perfil_paciente

    f_a = ficha(org_a, admin_a, practitioner_a, branches_a["centro"], agenda_a, 5, status="attended")
    pagar(org_a, f_a, admin_a, "100.00")

    perfil_paciente(user_b, org_b)
    dar_rol(user_b, org_b, "admin-operativo", "Administración",
            ["reporting.kpi.read", "reporting.report.run"])
    with tenant_context(org_b.id):
        prof_b = Practitioner.objects.create(
            organization=org_b, first_name="Bruno", last_name="Paz",
            license_number="MP-3000",
        )
        PractitionerBranch.objects.create(
            organization=org_b, practitioner=prof_b, branch=branch_b,
        )
        agenda_b = Schedule.objects.create(
            organization=org_b, practitioner=prof_b, branch=branch_b,
            weekday=agenda_a.weekday, start_time="09:00", end_time="11:00",
            slot_minutes=60, valid_from=agenda_a.valid_from,
        )
    for dias, monto in ((9, "300.00"), (8, "200.00")):
        f_b = ficha(org_b, user_b, prof_b, branch_b, agenda_b, dias, status="attended")
        pagar(org_b, f_b, user_b, monto)

    cliente = autenticar(api_client, admin_a)
    assert len(correr(cliente, dataset="appointments").data["rows"]) == 1
    assert len(correr(cliente, dataset="payments").data["rows"]) == 1

    cliente = autenticar(api_client, user_b)
    assert len(correr(cliente, dataset="appointments").data["rows"]) == 2
    assert len(correr(cliente, dataset="payments").data["rows"]) == 2
